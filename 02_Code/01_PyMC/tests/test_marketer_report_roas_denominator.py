"""The displayed return must use the same budget as its denominator label."""

from __future__ import annotations

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

REPORT_PATH = Path(__file__).resolve().parents[2] / "02_Budget_optimizer" / "marketer_report.py"
SPEC = importlib.util.spec_from_file_location("marketer_report_roas_contract", REPORT_PATH)
assert SPEC is not None and SPEC.loader is not None
REPORT = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = REPORT
SPEC.loader.exec_module(REPORT)


class ReportRoasDenominatorTests(unittest.TestCase):
    @staticmethod
    def _inputs() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
        common = {
            "campaign_name": "TCX", "requested_budget_mln_rub": 100.0,
            "rto_optimizer_use": "no_increase:2", "quality_explanation": "Synthetic fixture",
            "strong_support_warnings_n": 0, "policy_violations_n": 0,
        }
        scenarios = pd.DataFrame([
            {**common, "scenario_no": "S01", "scenario_name": "Исходный полный план",
             "candidate_name": "TCX__scenario1_current_plan", "budget_mln_rub": 100.0,
             "allocated_budget_mln_rub": 100.0, "unallocated_budget_mln_rub": 0.0,
             "allocated_budget_share": 1.0, "rto_p10_mln": 100.0,
             "rto_p50_mln": 150.0, "rto_p90_mln": 200.0,
             "rto_roas_p50": 1.5, "roas_allocated_budget_p50": 1.5,
             "roas_requested_budget_p50": 1.5,
             "roas_denominator_kind": "requested_budget",
             "roas_denominator_budget_rub": 100_000_000.0,
             "hard_support_warnings_n": 1, "quality_status": "Требуется ручная проверка",
             "scenario_variant": "uploaded_plan", "scenario_feasibility_status": "feasible_full"},
            {**common, "scenario_no": "S05", "scenario_name": "Безопасно распределяемая часть",
             "candidate_name": "TCX__scenario5_safe_partial", "budget_mln_rub": 80.0,
             "allocated_budget_mln_rub": 80.0, "unallocated_budget_mln_rub": 20.0,
             "allocated_budget_share": 0.8, "rto_p10_mln": 70.0,
             "rto_p50_mln": 96.0, "rto_p90_mln": 120.0,
             "rto_roas_p50": 1.2, "roas_allocated_budget_p50": 1.2,
             "roas_requested_budget_p50": 0.96,
             "roas_denominator_kind": "allocated_budget",
             "roas_denominator_budget_rub": 80_000_000.0,
             "hard_support_warnings_n": 0, "quality_status": "Сопоставимо с историей",
             "scenario_variant": "safe_partial", "scenario_feasibility_status": "feasible_partial"},
        ])
        context = pd.DataFrame([{
            "campaign_name": "TCX", "uploaded_budget_mln_rub": 100.0,
            "model_input_budget_mln_rub": 100.0, "unmodeled_budget_mln_rub": 0.0,
            "campaign_start": "2026-08-25", "campaign_end": "2026-09-21", "geos_n": 2,
        }])
        scenario6 = pd.DataFrame(columns=["campaign_name"])
        allocation = pd.DataFrame(columns=["source_campaign_name", "candidate_name"])
        return scenarios, scenario6, context, allocation

    def _pool_and_recommendations(self):
        scenarios, scenario6, context, allocation = self._inputs()
        before = scenarios.copy(deep=True)
        policy = {"reliability": {"usable_partial_coverage_min": 0.7}}
        pool = REPORT._build_decision_pool(scenarios, scenario6, context, allocation, policy)
        recommendations = REPORT._recommendations(
            scenarios, scenario6, context, min_roas_p50=1.0,
            allocation=allocation, decision_pool=pool, decision_policy=policy,
        )
        pd.testing.assert_frame_equal(before, scenarios)
        return pool, recommendations, scenario6, context, allocation

    def test_partial_and_full_primary_roas_match_the_declared_denominator(self) -> None:
        pool, recommendations, *_ = self._pool_and_recommendations()
        for _, row in pool.iterrows():
            with self.subTest(scenario=row["scenario_no"]):
                expected = row["rto_p50_mln"] * 1_000_000.0 / row["roas_denominator_budget_rub"]
                self.assertAlmostEqual(row["roas_p50"], expected)
        partial = pool[pool["scenario_no"].eq("S05")].iloc[0]
        self.assertEqual(partial["roas_denominator_kind"], "allocated_budget")
        self.assertAlmostEqual(partial["roas_p50"], 1.2)
        self.assertAlmostEqual(partial["roas_requested_budget_p50"], 0.96)
        chosen = recommendations.iloc[0]
        self.assertEqual(chosen["candidate_name"], "TCX__scenario5_safe_partial")
        self.assertEqual(chosen["decision_status"], "no_safe_recommendation")
        self.assertEqual(chosen["business_decision_status"], "Требуется ручное бизнес-решение")
        self.assertAlmostEqual(chosen["roas_p50"], 1.2)

    def test_legacy_primary_roas_survives_missing_or_unavailable_split_metrics(self) -> None:
        scenarios, scenario6, context, allocation = self._inputs()
        for missing_kind in (False, True):
            with self.subTest(missing_kind=missing_kind):
                legacy = scenarios.drop(columns=["roas_allocated_budget_p50", "roas_requested_budget_p50"])
                if missing_kind:
                    legacy = legacy.drop(columns=["roas_denominator_kind", "roas_denominator_budget_rub"])
                else:
                    legacy["roas_allocated_budget_p50"] = np.nan
                    legacy["roas_requested_budget_p50"] = np.nan
                pool = REPORT._build_decision_pool(legacy, scenario6, context, allocation, {})
                self.assertEqual(pool["roas_p50"].tolist(), [1.5, 1.2])

    def test_business_threshold_still_uses_requested_budget_roas(self) -> None:
        scenarios, scenario6, context, allocation = self._inputs()
        full = scenarios.iloc[:1].copy()
        full["hard_support_warnings_n"] = 0
        full["quality_status"] = "Сопоставимо с историей"
        pool = REPORT._build_decision_pool(full, scenario6, context, allocation, {})
        # Isolate the business gate from presentation: its input remains the
        # explicit requested-budget metric even if the displayed metric differs.
        pool["roas_p50"] = 1.2
        pool["roas_requested_budget_p50"] = 0.8
        chosen = REPORT._recommendations(
            full, scenario6, context, min_roas_p50=1.0, decision_pool=pool,
        ).iloc[0]
        self.assertEqual(chosen["candidate_name"], "TCX__scenario1_current_plan")
        self.assertEqual(chosen["business_decision_status"], "Ниже бизнес-порога")

    def test_workbook_main_conclusion_and_scenarios_show_the_actual_denominator(self) -> None:
        pool, recommendations, scenario6, context, allocation = self._pool_and_recommendations()
        label = "Бюджет в знаменателе ROAS, млн руб."
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            paths = REPORT.ReportPaths(root, root / "daily.csv", root, root / "report.xlsx", "fixture")
            REPORT._write_dynamic_workbook(
                paths, campaign_summary=context, scenario6=scenario6,
                recommendations=recommendations, decision_pool=pool, allocation=allocation,
                model_activation_status="preprod_restricted", production_blockers=[],
            )
            workbook = load_workbook(paths.output_xlsx, data_only=True)
            try:
                sheet = workbook["01_TCX"]
                for title in ("Главный вывод", "Все сценарии"):
                    with self.subTest(table=title):
                        title_row = next(cell.row for cell in sheet["A"] if cell.value == title)
                        header_row = title_row + 1
                        headers = [cell.value for cell in sheet[header_row]]
                        roas_col = headers.index("ROAS p50") + 1
                        denominator_col = headers.index(label) + 1
                        self.assertEqual(denominator_col, roas_col + 1)
                        values = []
                        for row in range(header_row + 1, sheet.max_row + 1):
                            if sheet.cell(row, 1).value is None:
                                break
                            values.append((sheet.cell(row, roas_col).value, sheet.cell(row, denominator_col).value))
                        expected = [(1.2, 80.0), (1.2, 80.0)] if title == "Главный вывод" else [(1.5, 100.0), (1.2, 80.0)]
                        self.assertEqual(values, expected)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
