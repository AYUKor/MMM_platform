"""Offline Excel projection using the application's openpyxl export stack."""
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

METHOD_TEXT = [
    "Модельная оценка эффекта кампании с учётом фактически проходившей параллельной рекламы",
    'Дополнительный РТО: фактическая реклама минус та же реклама без выбранной кампании.',
    'Параллельная реклама, ежедневные пользователи и покупки остаются фактическими.',
    'Размен = дополнительный РТО / полный исходный бюджет кампании.',
    'P10/P50/P90 считаются после суммирования каждого posterior draw. Квантили не складываются.',
    'Оценка условна на фактических пользователях; привлечение новых пользователей не оценивается.',
    'Модель имеет ограниченный исследовательский статус. Результат не является экспериментальным доказательством.',
    'Историческое распределение РФ сохранено; каталог будущих географий не используется.',
]


def export_report(path: Path, card: dict, daily: list, media: list) -> None:
    """Write numeric amounts and escape untrusted workbook text; reopen for acceptance."""
    book = Workbook()
    book.remove(book.active)
    summary = [['Показатель', 'P10', 'P50', 'P90'], ['Кампания', card['campaign_name']],
               ['Направление', card['segment']], ['Бюджет, руб.', card['source_budget_rub']]]
    for label, key in [('Дополнительный РТО, руб.', 'rto'), ('Размен', 'roas'),
                       ('РТО во время размещения, руб.', 'during'), ('РТО после размещения, руб.', 'after')]:
        summary.append([label] + [card['result'][key][q] for q in ['p10', 'p50', 'p90']])
    tables = {'Результат': summary,
              'По дням': [['Дата', 'Период', 'РТО P10', 'РТО P50', 'РТО P90']] +
                         [[r['date'], r['period'], *[r['rto'][q] for q in ['p10', 'p50', 'p90']]] for r in daily],
              'Фактический медиаплан': [['Дата', 'География', 'Медиаканал', 'Бюджет, руб.']] +
                                     [[r[k] for k in ['date', 'geography', 'media_channel', 'spend_rub']] for r in media],
              'Метод и ограничения': [['Методика']] + [[s] for s in METHOD_TEXT] +
                                     [['Версия модели', card['model_package_id']], ['Версия метода', card['method_version']]]}
    for title, rows in tables.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append([("'" + v if v.lstrip().startswith(('=', '+', '-', '@')) else v)
                          if isinstance(v, str) else v for v in row])
        sheet.freeze_panes = 'B2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='256347')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 28 if column[0].column > 1 else 48
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.000000'
    book.save(path)
    check = load_workbook(path, data_only=False, read_only=True)
    assert check.sheetnames == list(tables)
    assert abs(check['Результат']['B4'].value - card['source_budget_rub']) < 1e-7
    for i, key in enumerate(['rto', 'roas', 'during', 'after'], 5):
        for j, q in enumerate(['p10', 'p50', 'p90'], 2):
            value = check['Результат'].cell(i, j).value
            assert abs(value - card['result'][key][q]) <= max(1e-12, abs(value) * 1e-14)
    assert all(cell.data_type != 'f' for sheet in check for row in sheet for cell in row)
    check.close()
