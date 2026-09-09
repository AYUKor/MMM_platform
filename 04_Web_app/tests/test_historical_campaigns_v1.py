"""Synthetic public fixtures only; real numerical acceptance remains local."""
from __future__ import annotations
import argparse
import copy
import hashlib
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

WEB = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(WEB))
sys.path.insert(0, str(WEB.parent / '02_Code/01_PyMC'))
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from mmm_core.historical_campaigns import (HistoricalContext, assess_campaign, campaign_state,
                                          evaluate_campaign, independent_response, response)
from services.historical_campaign_batch import run, seal_files, write_json, verify_seal
from services.historical_campaign_dataset import HistoricalDatasetStore, HistoricalDatasetError, DIRECTIONS
from services.historical_campaign_report import export_report, METHOD_TEXT
from contracts.historical_campaigns_v1 import validate_historical_campaigns
from tests import test_product_navigation_http_v1 as navigation_http


def synthetic_card(key='synthetic_a', budget=1000.0, start='2025-01-03'):
    q = {'p10':.001, 'p50':.002, 'p90':.003}
    return {'campaign_key':key,'source_group_id':'source-a','business_campaign_id':key,
            'campaign_name':'Синтетическая кампания','segment':DIRECTIONS[0],
            'start_date':start,'end_date':'2025-01-04','evaluation_end':'2025-01-06',
            'declared_windows':['2025-01-03 / 2025-01-04'],'tail_days':2,'historical_geographies':['Синтетический город'],
            'source_budget_rub':budget,'allocated_budget_rub':budget,'panel_budget_rub':budget,
            'evaluated_budget_rub':budget,'has_federal_rows':False,'geographies_count':1,
            'identity_status':'confirmed','coverage_status':'full','result_status':'calculated',
            'identity_reasons':[],'coverage_reasons':[],'reason_texts':[],'status_text':'Рассчитано',
            'result':{**{k:copy.deepcopy(q) for k in ['rto','roas','during','after']},'draws':8},
            'model_package_id':'synthetic_package','method_version':'historical-paired-media-v1','limitations':METHOD_TEXT}


def synthetic_dataset(root: Path, cards=None):
    directory=root/'synthetic_v1'; directory.mkdir(parents=True)
    cards=cards or [synthetic_card()]
    write_json(directory/'serving/registry.json', cards)
    for card in cards:
        if card['result_status'] != 'calculated':
            continue
        p=directory/'serving'/card['campaign_key'];p.mkdir()
        daily=[{'date':'2025-01-03','period':'Размещение','rto':card['result']['rto']}]
        media=[{'date':'2025-01-03','geography':'Синтетический город','media_channel':'TV','spend_rub':1000.0}]
        write_json(p/'daily.json',{'items':daily})
        write_json(p/'media.json',{'plan':media,'geo_channel_totals':[], 'geography_totals':[], 'channel_totals':[]})
        export_report(p/'report.xlsx',card,daily,media)
    manifest={'schema_version':'1.0.0','dataset_id':'synthetic_v1','acceptance':'passed',
              'model_package_id':'synthetic_package','method_version':'historical-paired-media-v1',
              'files':seal_files(directory,[p for p in (directory/'serving').rglob('*') if p.is_file()])}
    write_json(directory/'manifest.json',manifest)
    return HistoricalDatasetStore(root,'synthetic_v1','another_active_model')


def synthetic_context():
    dates=pd.date_range('2025-01-01',periods=12)
    frame=pd.DataFrame({'date':dates,'geo_label':'synthetic_geo','network':'ТС5','channel':'Онлайн',
                        'row_position':range(12),'unique_users':10.,'population_k':2.,'spend_TV':30.})
    row=pd.Series({'campaign_key':'synthetic_a','campaign_id':'a','campaign_name':'Synthetic',
                   'segment':DIRECTIONS[0],'start_date':dates[3],'end_date':dates[4],
                   'source_budget_rub':20.,'panel_budget_rub':20.,'has_federal_rows':False,
                   'source_declared_windows':'["2025-01-04 / 2025-01-05"]','source_declared_window_count':1})
    spend=frame.iloc[3:5][['date','geo_label','network','channel']].assign(campaign_key='synthetic_a',media_channel='TV',spend_rub=10.)
    transform={'l_max':2,'channels':['TV'],'spend_active':['spend_TV'],'geos':['synthetic_geo'],
               'x_scale_geo':[[5.]],'y_scale':3.}
    draws={'pairs':np.array([[0,0],[0,1],[1,0],[1,1]]),'alpha':np.array([[.2],[.4],[.6],[.8]]),
           'lam':np.ones((4,1)),'beta':np.ones((4,1,1))*2,'beta_kind':'pooled','beta_labels':[]}
    config={'train_start':'2025-01-01','train_end':'2025-01-12'}
    fit=DIRECTIONS[0]+'::turnover_per_user'
    return HistoricalContext({}, {}, pd.DataFrame([row]), spend, {fit:frame}, {fit:transform}, {fit:draws},config)


class HistoricalCalculationTest(unittest.TestCase):
    def test_paired_formula_parallel_spend_and_draw_aggregation(self):
        ctx=synthetic_context();row=ctx.registry.iloc[0];fit=next(iter(ctx.frames));f=ctx.frames[fit];t=ctx.transforms[fit];d=ctx.draws[fit]
        result,dates,geo,during,after,checks=evaluate_campaign(row,ctx.spend,f,t,d,chunk_size=3)
        x=f.spend_TV.to_numpy()/10
        alternative=x.copy();alternative[3:5]-=1
        args=(d['alpha'][:,0],d['lam'][:,0],d['beta'][:,0,0],f.unique_users.to_numpy(),3.,2)
        expected=independent_response(x,*args)-independent_response(alternative,*args)
        np.testing.assert_allclose(result,expected[:,3:7],rtol=1e-12,atol=1e-12)
        np.testing.assert_allclose(during+after,result.sum(axis=1))
        self.assertEqual(checks['same_cell_parallel_rows'],2)
        self.assertEqual(checks['parallel_remaining_spend_rub'],40.)
        np.testing.assert_array_equal(expected[:,:3],0)
        np.testing.assert_array_equal(expected[:,7:],0)

    def test_identity_is_independent_of_coverage(self):
        ctx=synthetic_context();row=ctx.registry.iloc[0];fit=next(iter(ctx.frames))
        a=assess_campaign(row,ctx.spend,ctx.frames[fit],ctx.transforms[fit])
        self.assertTrue(a['eligible'])
        row=row.copy();row['source_declared_windows']='["first", "second"]';row['source_declared_window_count']=2
        state=campaign_state(row,a,ctx.config)
        self.assertEqual((state['identity_status'],state['coverage_status'],state['result_status']),('needs_review','full','unavailable'))
        ctx.config['train_end']='2025-01-04'
        state=campaign_state(row,a,ctx.config)
        self.assertIn('CROSSES_TRAINING_END',state['coverage_reasons'])
        self.assertIn('TAIL_UNAVAILABLE',state['coverage_reasons'])

    def test_mandatory_inputs_never_skip_missing_or_changed_file(self):
        with tempfile.TemporaryDirectory() as temp:
            p=Path(temp)/'mandatory';p.write_text('original')
            ctx=HistoricalContext({'files':{'posterior':{'sha256':hashlib.sha256(b'original').hexdigest()}}},{'posterior':p},None,None,{}, {},{}, {})
            ctx.verify();p.write_text('changed')
            with self.assertRaises(ValueError):ctx.verify()
            p.unlink()
            with self.assertRaises(ValueError):ctx.verify()
            with self.assertRaises(ValueError):HistoricalContext.load({'files':{}},{},Path(temp))

    def test_physically_complete_inputs_and_each_mandatory_file(self):
        from tests.synthetic_historical_inputs import write_inputs
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);inputs=write_inputs(root)
            ctx=HistoricalContext.load(inputs['spec'],inputs['roots'],root/'private')
            self.assertEqual(len(ctx.frames),4)
            self.assertEqual({len(d['pairs']) for d in ctx.draws.values()},{4})
            for role,path in ctx.paths.items():
                original=path.read_bytes()
                path.write_bytes(b'corrupt')
                with self.subTest(role=role),self.assertRaises(ValueError):ctx.verify()
                path.unlink()
                with self.subTest(missing=role),self.assertRaises(ValueError):ctx.verify()
                path.write_bytes(original)
            args=argparse.Namespace(input_manifest=str(root/'input.json'),output_root=str(root/'datasets'),dataset_id='physical',
                                    package_root=str(inputs['roots']['package']),data_root=str(inputs['roots']['data']),
                                    source_root=str(inputs['roots']['source']),resume=False,max_campaigns=0)
            outcome=run(args)
            self.assertEqual(sum(c['calculated'] for c in outcome['counts'].values()),4)
            from services.historical_campaign_dataset import verify_materialized
            self.assertEqual(verify_materialized(root/'datasets','physical')['campaigns'],4)

    def test_global_corruption_prevents_publication_but_campaign_error_isolated(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);spec=root/'input.json'
            write_json(spec,{'model_package_id':'synthetic','package_fingerprint':'a'*64,'files':{'panel':{'sha256':'b'*64}},'goldens':{}})
            args=argparse.Namespace(input_manifest=str(spec),output_root=str(root/'datasets'),dataset_id='broken',package_root=temp,data_root=temp,source_root=temp,resume=False,max_campaigns=0)
            ctx=synthetic_context()
            with patch.object(HistoricalContext,'load',return_value=ctx),patch.object(ctx,'verify',side_effect=ValueError('changed input')):
                with self.assertRaises(ValueError):run(args)
            self.assertFalse((root/'datasets/broken/manifest.json').exists())
            args.dataset_id='isolated'
            other=ctx.registry.copy();other['campaign_key']='synthetic_b';ctx.registry=pd.concat([ctx.registry,other],ignore_index=True)
            other_spend=ctx.spend.copy();other_spend['campaign_key']='synthetic_b';ctx.spend=pd.concat([ctx.spend,other_spend])
            def fail_one(row,*rest,**kwargs):
                if row.campaign_key=='synthetic_b':raise ValueError('campaign-specific invalid input')
                return evaluate_campaign(row,*rest,**kwargs)
            with patch.object(HistoricalContext,'load',return_value=ctx),patch.object(ctx,'verify'),patch('services.historical_campaign_batch.evaluate_campaign',side_effect=fail_one):
                result=run(args)
            self.assertEqual(result['counts'][DIRECTIONS[0]]['calculated'],1)
            self.assertEqual(result['counts'][DIRECTIONS[0]]['errors'],1)
            store=HistoricalDatasetStore(root/'datasets','isolated')
            self.assertEqual(store.read({},'synthetic_b','card')['campaign']['result_status'],'error')

    def test_resume_preserves_completed_result_and_rejects_drift(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);spec=root/'input.json'
            write_json(spec,{'model_package_id':'synthetic','package_fingerprint':'a'*64,'files':{'panel':{'sha256':'b'*64}},'goldens':{}})
            args=argparse.Namespace(input_manifest=str(spec),output_root=str(root/'datasets'),dataset_id='resume_test',package_root=temp,data_root=temp,source_root=temp,resume=False,max_campaigns=1)
            ctx=synthetic_context()
            # Two campaigns, disjoint IDs; one result per invocation, actual calculation.
            second=ctx.registry.copy();second['campaign_key']='synthetic_b';ctx.registry=pd.concat([ctx.registry,second],ignore_index=True)
            second_spend=ctx.spend.copy();second_spend['campaign_key']='synthetic_b';ctx.spend=pd.concat([ctx.spend,second_spend])
            with patch.object(HistoricalContext,'load',return_value=ctx),patch.object(ctx,'verify'):
                first=run(args);self.assertEqual(first['pending'],1)
                progress=json.loads((root/'datasets/resume_test/progress.json').read_text())
                original=progress['completed']['synthetic_a']['files']
                args.resume=True;second=run(args);self.assertEqual(second['pending'],0)
                verify_seal(root/'datasets/resume_test',original)
                with self.assertRaises(ValueError):run(args)


class HistoricalStoreTest(unittest.TestCase):
    def test_search_before_pagination_numeric_sort_ties_and_small_ratio(self):
        with tempfile.TemporaryDirectory() as temp:
            cards=[synthetic_card('synthetic_a',9),synthetic_card('synthetic_b',100),synthetic_card('synthetic_c',9)]
            cards[0]['campaign_name']='Другое';cards[1]['campaign_name']=' Поиск ';cards[2]['campaign_name']='ПОИСК'
            store=synthetic_dataset(Path(temp),cards)
            result=store.read({'q':['поиск'],'sort':['source_budget_rub'],'order':['asc'],'limit':['1']})
            self.assertEqual(len(result['blocks']),4)
            self.assertEqual(result['blocks'][0]['total'],2)
            self.assertEqual(result['blocks'][0]['items'][0]['campaign_key'],'synthetic_c')
            self.assertEqual(result['blocks'][0]['items'][0]['result']['roas']['p50'],.002)
            validate_historical_campaigns(result)
            for kind in ['card','daily','media']:
                validate_historical_campaigns(store.read({},'synthetic_a',kind))
            self.assertFalse(result['matches_active_model'])
            payload=store.read({},'synthetic_a','report.xlsx')
            wb=load_workbook(io.BytesIO(payload),data_only=True)
            self.assertEqual(wb['Результат']['C6'].value,.002)
            self.assertEqual(wb['Результат']['B4'].value,9)
            wb.close()

    def test_six_sorts_null_last_stable_keys_and_separate_directions(self):
        with tempfile.TemporaryDirectory() as temp:
            cards=[synthetic_card('synthetic_b',9,'2025-01-03'), synthetic_card('synthetic_a',9,'2025-01-03'),
                   synthetic_card('synthetic_c',100,'2025-01-04'), synthetic_card('synthetic_null'), synthetic_card('other_direction')]
            cards[2]['end_date']='2025-01-05'
            cards[3].update(start_date=None,end_date=None,source_budget_rub=None,result=None,
                            result_status='unavailable',coverage_status='unavailable')
            cards[4]['segment']=DIRECTIONS[1]
            store=synthetic_dataset(Path(temp),cards)
            for sort in ['start_date','end_date','source_budget_rub']:
                for order in ['asc','desc']:
                    result=store.read({'sort':[sort],'order':[order],'limit':['100']})
                    keys=[r['campaign_key'] for r in result['blocks'][0]['items']]
                    expected=['synthetic_a','synthetic_b','synthetic_c'] if order=='asc' else ['synthetic_c','synthetic_a','synthetic_b']
                    self.assertEqual(keys,expected+['synthetic_null'])
                    self.assertEqual(result['blocks'][1]['total'],1)
                    self.assertIsNone(result['blocks'][0]['items'][-1]['result'])
                    paged=store.read({'sort':[sort],'order':[order],'offset':['1'],'limit':['1']})
                    self.assertEqual(paged['blocks'][0]['items'][0]['campaign_key'],expected[1])
            with self.assertRaises(HistoricalDatasetError) as error:store.read({},'synthetic_null','daily')
            self.assertEqual(error.exception.code,'HISTORICAL_RESULT_UNAVAILABLE')
            manifest_path=Path(temp)/'synthetic_v1/manifest.json'
            manifest=json.loads(manifest_path.read_text());manifest['schema_version']='99.0.0';write_json(manifest_path,manifest)
            with self.assertRaises(HistoricalDatasetError) as error:store.read({})
            self.assertEqual(error.exception.code,'HISTORICAL_VERSION_UNAVAILABLE')

    def test_persistence_fresh_process_without_batch_imports(self):
        with tempfile.TemporaryDirectory() as temp:
            synthetic_dataset(Path(temp))
            code="import sys,json;from pathlib import Path;sys.path.insert(0,sys.argv[1]);from services.historical_campaign_dataset import HistoricalDatasetStore;s=HistoricalDatasetStore(Path(sys.argv[2]),'synthetic_v1');print(json.dumps(s.read({},'synthetic_a','card')));assert 'mmm_core.historical_campaigns' not in sys.modules"
            result=subprocess.run([sys.executable,'-B','-c',code,str(WEB),temp],check=True,capture_output=True,text=True)
            self.assertEqual(json.loads(result.stdout)['campaign']['result']['roas']['p50'],.002)

    def test_query_unknown_version_and_corrupt_are_distinct(self):
        with tempfile.TemporaryDirectory() as temp:
            store=synthetic_dataset(Path(temp))
            for params,key,resource,code in [({'limit':['0']},None,'registry','QUERY_INVALID'),({},'missing','card','NOT_FOUND'),({'dataset_id':['missing']},None,'registry','VERSION_UNAVAILABLE'),({'dataset_id':['../escape']},None,'registry','QUERY_INVALID')]:
                with self.assertRaises(HistoricalDatasetError) as caught:store.read(params,key,resource)
                self.assertEqual(caught.exception.code,'HISTORICAL_'+code)
            (Path(temp)/'synthetic_v1/serving/registry.json').write_text('[]')
            with self.assertRaises(HistoricalDatasetError) as caught:store.read({})
            self.assertEqual(caught.exception.code,'HISTORICAL_DATASET_CORRUPT')


class HistoricalHttpTest(unittest.TestCase):
    tearDown = navigation_http.ProductNavigationHttpTest.tearDown
    request = navigation_http.ProductNavigationHttpTest.request
    def setUp(self):
        navigation_http.ProductNavigationHttpTest.setUp(self)
        self.application.historical=synthetic_dataset(Path(self.temporary.name)/'history')

    def test_historical_routes_use_saved_results(self):
        for path in ['', '/synthetic_a','/synthetic_a/daily','/synthetic_a/media']:
            status,payload=self.request('/api/v1/historical-campaigns'+path)
            self.assertEqual(status,200)
            validate_historical_campaigns(payload)
        self.assertEqual(self.request('/api/v1/historical-campaigns?limit=0')[0],422)
        self.assertEqual(self.request('/api/v1/historical-campaigns/missing')[0],404)
        self.assertEqual(self.request('/api/v1/historical-campaigns?dataset_id=missing')[0],404)
        import urllib.request
        request=urllib.request.Request(self.base_url+'/api/v1/historical-campaigns/synthetic_a/report.xlsx',headers={'Cookie':self.session_cookie})
        with urllib.request.urlopen(request) as response:
            self.assertIn('attachment',response.headers['Content-Disposition'])
            wb=load_workbook(io.BytesIO(response.read()),data_only=True)
            self.assertEqual(wb['Результат']['C6'].value,.002)
            wb.close()
        from dataclasses import replace
        provider=self.application.auth.identity_provider
        original=provider.resolve_session
        def restricted(token, *, request_id):
            resolution=original(token,request_id=request_id)
            return replace(resolution,context=replace(resolution.context,permissions=()))
        with patch.object(provider,'resolve_session',side_effect=restricted):
            for path in ['', '/synthetic_a/report.xlsx']:
                self.assertEqual(self.request('/api/v1/historical-campaigns'+path)[0],403)
        self.session_cookie=''
        self.assertEqual(self.request('/api/v1/historical-campaigns')[0],401)


if __name__=='__main__':unittest.main()
